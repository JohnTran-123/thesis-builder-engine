#!/usr/bin/env python3
"""
build_thesis.py — generic markdown-to-docx thesis assembler.

Assembles a thesis .docx from chapter drafts in markdown, using a
school-supplied template for fonts, spacing, and heading styles.

This is a v1 generic builder. For institution-specific features
(footers, page numbers, ToC fields, table captions), customize
the script after inspecting your school template's structure.

Usage:
    python build_thesis.py --template path/to/school_template.docx

Defaults assume the standard engine workspace layout:
    - Drafts in workspace/drafts/
    - Output to workspace/final/thesis.docx
    - Optional bibliography file (extracted from workspace/plan/plan.md by /presubmit)
"""

import argparse
import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("ERROR: python-docx not installed. Run: pip install python-docx", file=sys.stderr)
    sys.exit(1)


def parse_args():
    parser = argparse.ArgumentParser(description="Generic markdown-to-docx thesis builder")
    parser.add_argument("--drafts", default="workspace/drafts",
                        help="Directory of chapter drafts (default: workspace/drafts)")
    parser.add_argument("--template", required=True,
                        help="School-supplied docx template (REQUIRED)")
    parser.add_argument("--output", default="workspace/final/thesis.docx",
                        help="Output docx path (default: workspace/final/thesis.docx)")
    parser.add_argument("--order",
                        help="Optional text file listing chapter filenames in order; one filename per line. If omitted, files are assembled alphabetically.")
    parser.add_argument("--bibliography",
                        help="Optional bibliography file path. If provided, content is appended under a 'References' heading.")
    return parser.parse_args()


def load_chapter_order(args):
    """Return ordered list of chapter file paths."""
    drafts_dir = Path(args.drafts)
    if not drafts_dir.is_dir():
        print(f"ERROR: drafts dir not found: {drafts_dir}", file=sys.stderr)
        sys.exit(1)

    if args.order:
        order_path = Path(args.order)
        if not order_path.is_file():
            print(f"ERROR: order file not found: {order_path}", file=sys.stderr)
            sys.exit(1)
        with open(order_path, encoding="utf-8") as f:
            filenames = [line.strip() for line in f if line.strip() and not line.startswith("#")]
        return [drafts_dir / fn for fn in filenames]

    return sorted(drafts_dir.glob("*.md"))


def md_to_paragraphs(doc, md_text):
    """Append markdown content to doc as paragraphs.

    Minimal markdown-to-docx conversion (v1):
    - '# ' through '#### ' map to Heading 1-4
    - Blank lines map to empty paragraphs
    - Everything else is body paragraph text
    - Inline markdown (**bold**, *italic*, links) is NOT yet rendered
    - Lists, code blocks, tables are emitted as body text without
      special formatting (extend the script if you need richer support)
    """
    for line in md_text.split("\n"):
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=4)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.strip() == "":
            doc.add_paragraph("")
        else:
            doc.add_paragraph(line)


def append_bibliography(doc, bib_path):
    """Append a 'References' heading + bibliography content."""
    if not bib_path.is_file():
        print(f"WARNING: bibliography not found: {bib_path}", file=sys.stderr)
        return
    doc.add_heading("References", level=1)
    with open(bib_path, encoding="utf-8") as f:
        for line in f:
            doc.add_paragraph(line.rstrip("\n"))


def main():
    args = parse_args()

    template_path = Path(args.template)
    if not template_path.is_file():
        print(f"ERROR: template not found: {template_path}", file=sys.stderr)
        sys.exit(1)

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    chapter_paths = load_chapter_order(args)
    if not chapter_paths:
        print(f"ERROR: no .md files found in {args.drafts}", file=sys.stderr)
        sys.exit(1)

    print(f"Template: {template_path}")
    print(f"Output:   {output_path}")
    print(f"Drafts:   {len(chapter_paths)} chapter(s)")

    doc = Document(str(template_path))

    for chapter_path in chapter_paths:
        print(f"  + {chapter_path.name}")
        if not chapter_path.is_file():
            print(f"    WARNING: file listed in order but missing: {chapter_path}", file=sys.stderr)
            continue
        with open(chapter_path, encoding="utf-8") as f:
            md_to_paragraphs(doc, f.read())

    if args.bibliography:
        append_bibliography(doc, Path(args.bibliography))

    doc.save(str(output_path))
    print(f"\nBuilt: {output_path}")


if __name__ == "__main__":
    main()
